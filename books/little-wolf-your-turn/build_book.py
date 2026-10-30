#!/usr/bin/env python3
"""Build editable book layouts and embedded-font PDFs from one checked source."""
import argparse
from collections import Counter
from functools import lru_cache
import html
import io
import json
import os
from pathlib import Path
import re
import xml.etree.ElementTree as ET

import fitz
from fontTools.ttLib import TTFont as FontToolsFont
from PIL import Image, ImageDraw
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Pt
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.pdfgen import canvas
from reportlab.lib.utils import ImageReader

ROOT = Path(__file__).resolve().parent
BOOK = json.loads((ROOT / 'story.json').read_text())
W, H = 960, 540
INK, MUTED = '#303536', '#626B68'
FONT = ROOT / 'assets/fonts/Iansui-Regular.ttf'
TARGETS = set(BOOK['coreCharacters'] + BOOK['extensionCharacters'])
STORY_HAN = len(re.findall(r'[\u3400-\u9fff]', ''.join(''.join(s['lines']) for s in BOOK['spreads'])))
pdfmetrics.registerFont(TTFont('Iansui', str(FONT)))
CROPS = {'wolf': (0, 0, .337, .49), 'dog': (.337, 0, .667, .49),
         'fox': (.667, 0, 1, .49), 'monkey': (0, .49, .337, 1),
         'lion': (.337, .49, .667, 1)}
HEAD_CROPS = {'wolf': (.04, 0, .31, .30), 'dog': (.35, .01, .65, .30),
              'fox': (.67, 0, .98, .31), 'monkey': (.04, .53, .30, .81),
              'lion': (.34, .48, .63, .82)}


def text_item(text, x, y, size=23, colour=INK, emphasis=False):
    sizes = [size * (1.12 if emphasis and ch in TARGETS else 1) for ch in text]
    advances = [pdfmetrics.stringWidth(ch, 'Iansui', fs) for ch, fs in zip(text, sizes)]
    return dict(kind='text', text=text, x=x, y=y, size=size, colour=colour,
                sizes=sizes, advances=advances, width=sum(advances), height=max(sizes, default=size) * 1.25)


def wrap(text, width, size, emphasis):
    closing = set('，。！？：；、」』）》】…')
    opening = set('「『（《【')
    lines, line = [], ''
    for ch in text:
        candidate = line + ch
        if line and text_item(candidate, 0, 0, size, emphasis=emphasis)['width'] > width and ch not in closing:
            if line[-1] in opening:
                carry, line = line[-1], line[:-1]
            else:
                carry = ''
            lines.append(line)
            line = carry + ch
        else:
            line = candidate
    if line:
        lines.append(line)
    return lines


def add_block(items, paragraphs, x, y, width=364, size=22.8, leading=33, colour=INK, emphasis=True):
    for para in paragraphs:
        fs = 33 if para == '「嗷——」' else size
        for line in wrap(para, width, fs, emphasis):
            items.append(text_item(line, x, y, fs, colour, emphasis))
            y += max(leading, fs * 1.25)
    return y


def picture(path, x, y, width, height, crop=None, alt='原創繪本插畫'):
    return dict(kind='image', path=str(path), x=x, y=y, width=width, height=height, crop=crop, alt=alt)


def rect(x, y, width, height, fill, stroke=None, radius=0):
    return dict(kind='rect', x=x, y=y, width=width, height=height, fill=fill, stroke=stroke, radius=radius)


def character(name, x, y, width, height, head=False):
    return picture(ROOT / 'assets/characters.png', x, y, width, height, (HEAD_CROPS if head else CROPS)[name], BOOK['characters'][name])


def pages(proof):
    result = []
    cover = ROOT / 'assets/illustrations/cover.png'
    if not proof and not cover.exists():
        raise FileNotFoundError('Missing finished cover illustration')
    front = [rect(240, 0, 480, 540, '#E8F1E9')]
    if cover.exists():
        front.append(picture(cover, 290, 239, 380, 296, alt='小狼拿着紙星，用自己的聲音說故事'))
    else:
        front.append(character('wolf', 330, 240, 290, 280))
    front += [text_item('部首故事・犬', 277, 28, 17, '#4D7568'),
              text_item('小狼，', 276, 72, 62), text_item('換你說！', 276, 143, 62)]
    result.append(dict(id='cover', items=front, note='原創封面。適讀年齡：7–9 歲；繁體中文犬部故事。'))
    title = [character('wolf', 94, 219, 290, 275),
             text_item('給每一個', 66, 85, 20), text_item('還沒說出願望的人。', 66, 117, 20),
             text_item('小狼，', 535, 167, 53), text_item('換你說！', 535, 232, 53),
             text_item('部首故事・犬', 539, 324, 19, '#4D7568')]
    title.append(text_item('原創犬部故事・7–9 歲', 539, 452, 13, MUTED))
    result.append(dict(id='title', items=title, note='右頁為扉頁 1，左頁是閱讀版的前襯畫面，不納入單頁 PDF 的內頁頁碼。'))
    for spread in BOOK['spreads']:
        art = ROOT / 'assets/illustrations' / (spread['id'] + '.png')
        items = []
        if art.exists():
            # Whitespace-rich plates allow deliberate movement without redrawing the art.
            x, y, width, height = (0, 0, W, H)
            if spread['id'] == 's02':
                x = -22
            elif spread['id'] == 's14':
                x, y, width, height = (0, 25, 900, 506.25)
            elif spread['id'] == 's10':
                y = 30
            elif spread['id'] == 's03':
                x, y, width, height = (170, 48, 790, 444.375)
            elif spread['layout'] == 'left-middle' and spread['id'] != 's04':
                x, y, width, height = (100, 28, 860, 483.75)
            items.append(picture(art, x, y, width, height, alt=spread['art']))
        elif proof:
            items += [rect(500, 205, 370, 240, '#F3F6F4'), text_item('插畫製作中', 594, 306, 22, MUTED)]
        else:
            raise FileNotFoundError('Missing finished illustration: ' + str(art))
        if spread['layout'] == 'left-middle':
            add_block(items, spread['lines'], 53, 154, 338, 22.8, 34)
        elif len(spread['lines']) == 2:
            add_block(items, spread['lines'], 52, 39)
        else:
            add_block(items, spread['lines'][:1], 52, 39)
            add_block(items, spread['lines'][1:], 532, 39)
        for number, x in zip(spread['pages'], (38, 905)):
            items.append(text_item(str(number), x, 513, 10.5, MUTED))
        result.append(dict(id=spread['id'], items=items,
                           note=f"內頁 {spread['pages'][0]}–{spread['pages'][1]}。{spread['beat']}\n" + '\n'.join(spread['lines'])))
    lesson = [text_item('字裏的朋友', 49, 39, 29), text_item('六個犬部字，五位故事朋友。', 49, 87, 17, MUTED)]
    cards = [('狗', '小狗', 'dog'), ('狼', '小狼', 'wolf'), ('猴', '猴子', 'monkey'),
             ('狐', '狐狸', 'fox'), ('狸', '狐狸', 'fox'), ('獅', '獅子', 'lion')]
    for index, (ch, word, name) in enumerate(cards):
        x, y = 48 + index % 3 * 130, 134 + index // 3 * 171
        lesson.append(rect(x, y, 112, 151, '#FFFFFF', '#D6E3DA', 9))
        lesson += [text_item(ch, x + 31, y + 6, 48), character(name, x + 24, y + 64, 64, 63, head=True), text_item(word, x + 38 - (len(word) - 2) * 7, y + 128, 15)]
    lesson += [text_item('犬和犭，是一家', 532, 39, 29),
               text_item('犬', 553, 104, 71, '#B35C50'), text_item('→', 658, 126, 34, MUTED),
               text_item('犭', 734, 104, 71, '#B35C50')]
    add_block(lesson, ['「犬」有狗的意思。', '它作左旁時，常寫成「犭」，', '這個形狀也叫反犬旁。'], 532, 221, 377, 18, 26, INK, False)
    lesson.append(text_item('猜一猜　　獨自', 532, 352, 28))
    add_block(lesson, ['「猜」和「獨」也是犬部字，', '卻不是動物的名稱。'], 532, 401, 377, 18, 26, INK, False)
    lesson += [text_item('再翻一遍，找出故事裏的犬部字。', 532, 477, 16, '#4D7568'),
               text_item('30', 38, 513, 10.5, MUTED), text_item('31', 905, 513, 10.5, MUTED)]
    result.append(dict(id='lesson', items=lesson, note='核心六字：狗、狐、狸、狼、猴、獅。延伸：猜、獨。詞語採用「狐狸」。繁體貓屬豸部，豬屬豕部，未誤列。'))
    back = [rect(480, 0, 480, 540, '#E8F1E9'),
            text_item('每個角色，', 54, 45, 28), text_item('都有自己的聲音。', 54, 85, 28),
            text_item('你想當甚麼？', 529, 105, 28), text_item('這一次，先聽你說。', 529, 151, 28),
            text_item('一個關於部首、角色', 529, 222, 18, '#4D7568'), text_item('與傾聽的故事。', 529, 250, 18, '#4D7568'),
            text_item('繁體中文・7–9 歲', 529, 497, 12, MUTED), text_item('32', 38, 513, 10.5, MUTED)]
    if cover.exists():
        back += [picture(cover, 95, 161, 228, 177), picture(cover, 668, 300, 190, 185)]
    else:
        back += [character('wolf', 110, 160, 220, 178), character('wolf', 668, 300, 190, 185)]
    add_block(back, ['《小狼，換你說！》', '原創故事與排版・點點書房', '插畫：AI 輔助，依原創角色設定製作', '字型：芫荽 Iansui・SIL OFL 1.1', f'正文 {STORY_HAN} 個漢字・六個核心字・兩個延伸字'], 54, 363, 382, 11.5, 20, MUTED, False)
    result.append(dict(id='back', items=back, note='左頁：製作頁 32；右頁：封底。閱讀後可請孩子設計自己的新角色卡。'))
    return result


def image_geometry(item):
    iw, ih = Image.open(item['path']).size
    if item['crop']:
        l, t, r, b = item['crop']
    else:
        l, t, r, b = 0, 0, 1, 1
    scale = min(item['width'] / ((r - l) * iw), item['height'] / ((b - t) * ih))
    width, height = (r - l) * iw * scale, (b - t) * ih * scale
    x = item['x'] + (item['width'] - width) / 2
    y = item['y'] + (item['height'] - height) / 2
    return dict(x=x, y=y, width=width, height=height,
                full_x=x - l * iw * scale, full_y=y - t * ih * scale,
                full_width=iw * scale, full_height=ih * scale, source=(iw, ih), crop=(l, t, r, b))


def colour(value):
    return tuple(int(value[i:i+2], 16) / 255 for i in (1, 3, 5))


@lru_cache(maxsize=20)
def pdf_image(path):
    image = Image.open(path)
    if 'A' in image.getbands():
        return ImageReader(path)
    # Standard high-quality PDF image encoding; the original PNG is kept intact.
    encoded = io.BytesIO()
    image.convert('RGB').save(encoded, format='JPEG', quality=95, subsampling=0, optimize=True)
    encoded.seek(0)
    return ImageReader(encoded)


def render_pdf(page_list, target):
    pdf = canvas.Canvas(str(target), pagesize=(W, H), pageCompression=1, initialFontName='Iansui')
    pdf.setTitle(BOOK['title'] + '｜犬部故事・跨頁閱讀版')
    pdf.setAuthor('點點書房・原創繪本製作')
    pdf.setSubject('7–9 歲繁體中文繪本；犬部與反犬旁；傾聽與自主表達')
    for page in page_list:
        for item in page['items']:
            kind = item['kind']
            if kind == 'rect':
                pdf.setFillColorRGB(*colour(item['fill']))
                pdf.setStrokeColorRGB(*colour(item['stroke'] or item['fill']))
                pdf.roundRect(item['x'], H - item['y'] - item['height'], item['width'], item['height'], item['radius'], stroke=int(bool(item['stroke'])), fill=1)
            elif kind == 'image':
                g = image_geometry(item)
                pdf.saveState()
                path = pdf.beginPath(); path.rect(g['x'], H - g['y'] - g['height'], g['width'], g['height'])
                pdf.clipPath(path, stroke=0)
                pdf.drawImage(pdf_image(item['path']), g['full_x'], H - g['full_y'] - g['full_height'], g['full_width'], g['full_height'], mask='auto')
                pdf.restoreState()
            else:
                pdf.setFillColorRGB(*colour(item['colour']))
                baseline = H - item['y'] - pdfmetrics.getAscent('Iansui') / 1000 * max(item['sizes'])
                obj = pdf.beginText(item['x'], baseline)
                for ch, fs in zip(item['text'], item['sizes']):
                    obj.setFont('Iansui', fs); obj.textOut(ch)
                pdf.drawText(obj)
        pdf.showPage()
    pdf.save()


def render_pptx(page_list, target):
    deck = Presentation(); deck.slide_width = Pt(W); deck.slide_height = Pt(H)
    deck.core_properties.title = BOOK['title']; deck.core_properties.subject = '犬部原創繪本・可編輯排版'
    for page in page_list:
        slide = deck.slides.add_slide(deck.slide_layouts[6])
        slide.notes_slide.notes_text_frame.text = page['note']
        for item in page['items']:
            if item['kind'] == 'rect':
                shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if item['radius'] else MSO_SHAPE.RECTANGLE,
                                               Pt(item['x']), Pt(item['y']), Pt(item['width']), Pt(item['height']))
                shape.fill.solid(); shape.fill.fore_color.rgb = RGBColor.from_string(item['fill'][1:])
                if item['stroke']:
                    shape.line.color.rgb = RGBColor.from_string(item['stroke'][1:]); shape.line.width = Pt(.6)
                else:
                    shape.line.fill.background()
            elif item['kind'] == 'image':
                g = image_geometry(item)
                shape = slide.shapes.add_picture(item['path'], Pt(g['x']), Pt(g['y']), Pt(g['width']), Pt(g['height']))
                l, t, r, b = g['crop']
                shape.crop_left = l; shape.crop_top = t; shape.crop_right = 1-r; shape.crop_bottom = 1-b
                shape._element.xpath('.//p:cNvPr')[0].set('descr', item['alt'])
            else:
                shape = slide.shapes.add_textbox(Pt(item['x']), Pt(item['y']), Pt(item['width'] + 5), Pt(item['height'] + 4))
                tf = shape.text_frame; tf.word_wrap = False
                tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
                p = tf.paragraphs[0]; p.space_before = p.space_after = Pt(0); p.line_spacing = 1.0
                for ch, fs in zip(item['text'], item['sizes']):
                    run = p.add_run(); run.text = ch; run.font.name = 'Iansui'; run.font.size = Pt(fs)
                    run.font.color.rgb = RGBColor.from_string(item['colour'][1:])
                    ea = OxmlElement('a:ea'); ea.set('typeface', 'Iansui'); run._r.get_or_add_rPr().append(ea)
    deck.save(str(target))


def render_svgs(page_list, directory):
    for page in page_list:
        elements = []
        for item in page['items']:
            if item['kind'] == 'rect':
                elements.append(f'<rect x="{item["x"]}" y="{item["y"]}" width="{item["width"]}" height="{item["height"]}" rx="{item["radius"]}" fill="{item["fill"]}" stroke="{item["stroke"] or "none"}"/>')
            elif item['kind'] == 'image':
                g = image_geometry(item); iw, ih = g['source']; l, t, r, b = g['crop']
                href = os.path.relpath(item['path'], directory)
                elements.append(f'<svg x="{g["x"]}" y="{g["y"]}" width="{g["width"]}" height="{g["height"]}" viewBox="{l*iw} {t*ih} {(r-l)*iw} {(b-t)*ih}"><title>{html.escape(item["alt"])}</title><image href="{html.escape(href)}" width="{iw}" height="{ih}"/></svg>')
            else:
                baseline = item['y'] + pdfmetrics.getAscent('Iansui') / 1000 * max(item['sizes'])
                spans = ''.join(f'<tspan font-size="{fs}">{html.escape(ch)}</tspan>' for ch, fs in zip(item['text'], item['sizes']))
                elements.append(f'<text x="{item["x"]}" y="{baseline}" fill="{item["colour"]}" font-family="Iansui,芫荽">{spans}</text>')
        content = '<svg xmlns="http://www.w3.org/2000/svg" width="960" height="540" viewBox="0 0 960 540"><rect width="960" height="540" fill="white"/>' + '\n'.join(elements) + '</svg>'
        (directory / (page['id'] + '.svg')).write_text(content)
        ET.fromstring(content)


def single_pages(source, target):
    spread = fitz.open(source); book = fitz.open()
    for index in range(len(spread)):
        if index == 0:
            clips = [fitz.Rect(240, 0, 720, H)]
        elif index == 1:
            clips = [fitz.Rect(W/2, 0, W, H)]
        else:
            clips = [fitz.Rect(0, 0, W/2, H), fitz.Rect(W/2, 0, W, H)]
        for clip in clips:
            page = book.new_page(width=W/2, height=H)
            page.show_pdf_page(page.rect, spread, index, clip=clip)
    book.set_metadata({'title': BOOK['title'] + '｜犬部故事', 'author': '點點書房・原創繪本製作', 'subject': '7–9 歲；封面、32內頁、封底'})
    book.set_toc([[1,'封面',1],[1,'故事',3],[1,'犬部遊戲',31],[1,'製作頁',33],[1,'封底',34]])
    book.save(target, garbage=4, deflate=True)


def validate(page_list, output, proof):
    cmap = FontToolsFont(FONT).getBestCmap()
    missing = sorted({ch for p in page_list for i in p['items'] if i['kind']=='text' for ch in i['text'] if ord(ch) not in cmap})
    assert not missing, 'Missing glyphs: ' + ''.join(missing)
    for p in page_list:
        for i in p['items']:
            if i['kind'] == 'text':
                assert i['x'] >= 0 and i['x'] + i['width'] <= W + .1 and i['y'] + i['height'] <= H, (p['id'], i['text'], 'outside page')
                if p['id'] != 'cover':
                    assert not (i['x'] < W/2 and i['x'] + i['width'] > W/2), (p['id'], i['text'], 'text crosses binding')
                else:
                    assert 240 <= i['x'] and i['x'] + i['width'] <= 720, (p['id'], 'text outside cover trim')
    spreads = fitz.open(output / '小狼，換你說！_跨頁閱讀版.pdf')
    leaves = fitz.open(output / '小狼，換你說！_繪本.pdf')
    assert len(spreads) == 18 and len(leaves) == 34
    leaf_index = 0
    for index, page in enumerate(page_list):
        if index == 0:
            clips = [(240,720)]
        elif index == 1:
            clips = [(480,960)]
        else:
            clips = [(0,480),(480,960)]
        for left,right in clips:
            expected = ''.join(i['text'] for i in page['items'] if i['kind']=='text' and i['x'] >= left and i['x'] + i['width'] <= right)
            actual = leaves[leaf_index].get_text(clip=leaves[leaf_index].rect)
            assert Counter(re.sub(r'\s','',expected)) == Counter(re.sub(r'\s','',actual)), ('Clipped or missing text',leaf_index+1)
            leaf_index += 1
    full_text = ''.join(p.get_text() for p in spreads)
    for ch in BOOK['coreCharacters'] + BOOK['extensionCharacters'] + ['犬','犭']:
        assert ch in full_text, 'PDF text missing: ' + ch
    deck = Presentation(output / '小狼，換你說！_可編輯排版.pptx')
    assert len(deck.slides) == 18
    assert sum(len(s.shapes) for s in deck.slides) > 100
    expected = ''.join(i['text'] for p in page_list for i in p['items'] if i['kind']=='text')
    actual = ''.join(shape.text for slide in deck.slides for shape in slide.shapes if shape.has_text_frame)
    assert expected == actual, 'Editable slide text differs from book source'
    report = {'status': 'proof' if proof else 'complete', 'spreadPages':len(spreads), 'singlePages':len(leaves),
              'editableSlides':len(deck.slides), 'missingGlyphs':missing, 'allTextEditable':True,
              'storyHanCharacters':len(re.findall(r'[\u3400-\u9fff]', ''.join(''.join(s['lines']) for s in BOOK['spreads']))),
              'targets':{ch:full_text.count(ch) for ch in BOOK['coreCharacters']+BOOK['extensionCharacters']}}
    (output / 'validation.json').write_text(json.dumps(report,ensure_ascii=False,indent=2)+'\n')
    return report


def main():
    parser = argparse.ArgumentParser(description=__doc__); parser.add_argument('--proof', action='store_true')
    args = parser.parse_args()
    output = ROOT / ('.cache/proof' if args.proof else 'exports')
    editable = output / 'editable' if args.proof else ROOT / 'editable'
    output.mkdir(parents=True, exist_ok=True); editable.mkdir(parents=True, exist_ok=True)
    page_list = pages(args.proof)
    render_pdf(page_list, output / '小狼，換你說！_跨頁閱讀版.pdf')
    single_pages(output / '小狼，換你說！_跨頁閱讀版.pdf', output / '小狼，換你說！_繪本.pdf')
    render_pptx(page_list, output / '小狼，換你說！_可編輯排版.pptx')
    render_svgs(page_list, editable)
    report = validate(page_list, output, args.proof)
    doc = fitz.open(output / '小狼，換你說！_跨頁閱讀版.pdf')
    preview = ROOT / '.cache' / ('proof/previews' if args.proof else 'final/previews')
    preview.mkdir(parents=True, exist_ok=True)
    thumbs = []
    for index, page in enumerate(doc):
        pix = page.get_pixmap(matrix=fitz.Matrix(1,1), alpha=False)
        pix.save(preview / f'{index+1:02}.png')
        thumb = Image.frombytes('RGB',[pix.width,pix.height],pix.samples); thumb.thumbnail((320,180))
        tile=Image.new('RGB',(336,208),'#e9ece9');tile.paste(thumb,(8,18));ImageDraw.Draw(tile).text((8,3),str(index+1),fill='#303536');thumbs.append(tile)
    sheet=Image.new('RGB',(336*3,208*6),'white')
    for index,thumb in enumerate(thumbs):sheet.paste(thumb,((index%3)*336,(index//3)*208))
    sheet.save(output / '全書縮圖.png')
    print(json.dumps(report,ensure_ascii=False,indent=2))


if __name__ == '__main__':
    main()
