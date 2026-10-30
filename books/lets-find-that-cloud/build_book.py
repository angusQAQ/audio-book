#!/usr/bin/env python3
"""Build editable book layouts and embedded-font PDFs from one checked source."""
import argparse
from collections import Counter
from functools import lru_cache
import html
import io
import json
import os
from pathlib import Path
import re
import subprocess
import sys
import xml.etree.ElementTree as ET

import fitz
from fontTools.ttLib import TTFont as FontToolsFont
from PIL import Image, ImageDraw
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Pt
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.pdfgen import canvas
from reportlab.lib.utils import ImageReader

ROOT = Path(__file__).resolve().parent
BOOK = json.loads((ROOT / 'story.json').read_text())
W, H = 960, 540
INK, MUTED = '#303536', '#626B68'
FONT = ROOT / 'assets/fonts/Iansui-Regular.ttf'
HAND_FONT = ROOT / 'assets/fonts/JyunsaiKaai-Bold.ttf'
HAND_FAMILY = 'JyunsaiKaai'
FONTS = {'Iansui': FONT, HAND_FAMILY: HAND_FONT}
EMPHASIS_SCALE = BOOK['typography']['targetScale']
TARGETS = set(BOOK['coreCharacters'] + BOOK['extensionCharacters'] + [BOOK['radical']['base']])
STORY_HAN = len(re.findall(r'[\u3400-\u9fff]', ''.join(''.join(s['lines']) for s in BOOK['spreads'])))
for family, path in FONTS.items():
    pdfmetrics.registerFont(TTFont(family, str(path)))
CROPS = {'ami': (0,0,.25,1), 'cheng': (.5,0,.75,1)}


def text_item(text, x, y, size=23, colour=INK, emphasis=False):
    fonts = [HAND_FAMILY if emphasis and ch in TARGETS else 'Iansui' for ch in text]
    sizes = [size * (EMPHASIS_SCALE if family == HAND_FAMILY else 1) for family in fonts]
    advances = [pdfmetrics.stringWidth(ch, family, fs) for ch, family, fs in zip(text, fonts, sizes)]
    # Share a baseline while allowing the larger handwritten glyphs to rise above the body.
    ascent = max((pdfmetrics.getAscent(family) / 1000 * fs for family, fs in zip(fonts, sizes)), default=size)
    descent = max((-pdfmetrics.getDescent(family) / 1000 * fs for family, fs in zip(fonts, sizes)), default=0)
    return dict(kind='text', text=text, x=x, y=y, size=size, colour=colour,
                fonts=fonts, sizes=sizes, advances=advances, ascent=ascent,
                width=sum(advances), height=ascent + descent)


def wrap(text, width, size, emphasis):
    closing, opening = set('，。！？：；、」』）》】…'), set('「『（《【')
    units = []
    for ch in text:
        if units and (ch in closing or units[-1][-1] in opening):
            units[-1] += ch
        else:
            units.append(ch)
    lines, line = [], ''
    for unit in units:
        if line and text_item(line+unit,0,0,size,emphasis=emphasis)['width'] > width:
            lines.append(line); line=unit
        else:
            line += unit
    if line: lines.append(line)
    return lines


def add_block(items, paragraphs, x, y, width=364, size=22.8, leading=33, colour=INK, emphasis=True):
    for para in paragraphs:
        fs = 33 if para == '「嗷——」' else size
        for line in wrap(para, width, fs, emphasis):
            item = text_item(line, x, y, fs, colour, emphasis)
            items.append(item)
            y += max(leading, item['height'] + 6)
    return y


def picture(path, x, y, width, height, crop=None, alt='原創繪本插畫'):
    return dict(kind='image', path=str(path), x=x, y=y, width=width, height=height, crop=crop, alt=alt)


def rect(x, y, width, height, fill, stroke=None, radius=0):
    return dict(kind='rect', x=x, y=y, width=width, height=height, fill=fill, stroke=stroke, radius=radius)


def character(name, x, y, width, height):
    return picture(ROOT/'assets/characters.png',x,y,width,height,CROPS[name],BOOK['characters'][name])


def pages(proof):
    result=[]
    cover=ROOT/'assets/illustrations/cover.png'
    if not proof and not cover.exists(): raise FileNotFoundError('Missing cover illustration')
    front=[rect(240,0,480,540,BOOK['palette']['cover']),
           text_item('部首故事・走',278,26,17,'#557D8B'),
           text_item('走！',278,65,56,emphasis=True),
           text_item('去找那朵雲',277,170,51)]
    if cover.exists():front.append(picture(cover,270,249,420,275,alt='阿米與小澄帶着紙風車，抬頭望雲'))
    elif proof:front += [character('ami',330,270,100,230),character('cheng',495,270,100,230)]
    result.append(dict(id='cover',items=front,note='原創走部繪本；封面；7–9 歲。'))
    title=[text_item('給每一個',63,78,21),text_item('走在路上的人。',63,113,21,emphasis=True),
           character('ami',75,248,125,242),character('cheng',247,248,125,242),
           text_item('走！',537,159,42,emphasis=True),text_item('去找那朵雲',536,242,46),
           text_item('部首故事・走',540,329,19),text_item('原創繁體中文繪本・7–9 歲',540,452,13,MUTED)]
    result.append(dict(id='title',items=title,note='右頁為扉頁 1；左頁為閱讀版的前襯，不計入單頁 PDF。'))
    for spread in BOOK['spreads']:
        items=[];art=ROOT/'assets/illustrations'/(spread['id']+'.png')
        if art.exists():items.append(picture(art,0,0,W,H,alt=spread['art']))
        elif proof:items.append(rect(515,265,378,211,'#F2F6F7'))
        else:raise FileNotFoundError('Missing illustration: '+str(art))
        if spread['layout']=='left-middle':
            add_block(items,spread['lines'],52,150,354,22.8,34)
        else:
            add_block(items,spread['lines'][:1],52,37,356,22.8,34)
            add_block(items,spread['lines'][1:],534,37,356,22.8,34)
        for number,x in zip(spread['pages'],(38,905)):items.append(text_item(str(number),x,513,10.5,MUTED))
        result.append(dict(id=spread['id'],items=items,note=f"內頁 {spread['pages'][0]}–{spread['pages'][1]}。{spread['beat']}\n"+'\n'.join(spread['lines'])))
    lesson=[text_item('字裏有一條小路',46,37,29),text_item('六個走部字，陪我們走一趟。',46,86,17,MUTED)]
    for n,(ch,word) in enumerate([('起','起風'),('趕','趕路'),('超','超過'),('越','越過'),('趟','一趟'),('趣','有趣')]):
        x,y=46+(n%3)*132,140+(n//3)*164
        lesson.append(rect(x,y,116,140,'#F4F8FA','#D3E0E6',8))
        lesson.append(text_item(ch,x+28,y+4,40,emphasis=True))
        example=text_item(word,0,0,18,emphasis=True);example['x']=x+(116-example['width'])/2;example['y']=y+98;lesson.append(example)
    lesson += [text_item('先看部首，再想意思',532,37,25),
               text_item('走',549,99,49,emphasis=True),text_item('走字旁',665,129,24,emphasis=True)]
    add_block(lesson,['「走」是這一家的部首。','放在左邊時，長長的捺筆','像小路一樣托住旁邊的部件。'],532,224,377,18,28)
    lesson.append(text_item('趁風小一點，慢慢走。',532,350,22,emphasis=True))
    add_block(lesson,['「趁」可以說把握合適的時候。'],532,400,377,18,28)
    lesson.append(text_item('跑、追、趕，哪個是走部字？',532,465,17,emphasis=True))
    lesson += [text_item('30',38,513,10.5,MUTED),text_item('31',905,513,10.5,MUTED)]
    result.append(dict(id='lesson',items=lesson,note='核心六字：起、趕、超、越、趟、趣；兼認走，延伸趁。答案：趕。跑屬足部，追屬辵部。'))
    back=[rect(480,0,480,540,BOOK['palette']['cover']),
          text_item('這一趟，',54,44,29,emphasis=True),text_item('你發現了甚麼？',54,94,29),
          text_item('雲在天上，',532,91,30),text_item('也在腳邊。',532,138,30),
          text_item('有趣的事，',532,216,23,emphasis=True),text_item('有時要慢一點才看得見。',532,258,23),
          text_item('繁體中文・7–9 歲',532,494,12,MUTED),text_item('32',38,513,10.5,MUTED)]
    if cover.exists():back.append(picture(cover,582,329,292,154,alt='兩位同行的孩子'))
    else:back += [character('ami',590,335,90,145),character('cheng',755,335,90,145)]
    back += [character('ami',91,167,98,170),character('cheng',258,167,98,170)]
    add_block(back,['《走！去找那朵雲》','原創故事與排版・點點書房','插畫：AI 輔助，依原創角色設定製作','字型：芫荽／芫茜雅楷・SIL OFL 1.1',f'正文 {STORY_HAN} 個漢字・六個核心字・一個延伸字','字卡答案：趕；跑屬足部，追屬辵部。'],54,367,382,11.5,19,MUTED,False)
    result.append(dict(id='back',items=back,note='左頁製作頁 32；右頁封底。快與慢都不是好壞的標籤，故事重點是注意與同行。'))
    return result


def image_geometry(item):
    iw, ih = Image.open(item['path']).size
    if item['crop']:
        l, t, r, b = item['crop']
    else:
        l, t, r, b = 0, 0, 1, 1
    scale = min(item['width'] / ((r - l) * iw), item['height'] / ((b - t) * ih))
    width, height = (r - l) * iw * scale, (b - t) * ih * scale
    x = item['x'] + (item['width'] - width) / 2
    y = item['y'] + (item['height'] - height) / 2
    return dict(x=x, y=y, width=width, height=height,
                full_x=x - l * iw * scale, full_y=y - t * ih * scale,
                full_width=iw * scale, full_height=ih * scale, source=(iw, ih), crop=(l, t, r, b))


def colour(value):
    return tuple(int(value[i:i+2], 16) / 255 for i in (1, 3, 5))


@lru_cache(maxsize=20)
def pdf_image(path):
    image = Image.open(path)
    if 'A' in image.getbands():
        return ImageReader(path)
    # Standard high-quality PDF image encoding; the original PNG is kept intact.
    encoded = io.BytesIO()
    image.convert('RGB').save(encoded, format='JPEG', quality=95, subsampling=0, optimize=True)
    encoded.seek(0)
    return ImageReader(encoded)


def render_pdf(page_list, target):
    pdf = canvas.Canvas(str(target), pagesize=(W, H), pageCompression=1, initialFontName='Iansui')
    pdf.setTitle(BOOK['title'] + '｜走部故事・跨頁閱讀版')
    pdf.setAuthor('點點書房・原創繪本製作')
    pdf.setSubject('7–9 歲繁體中文繪本；走部與走字旁；觀察、想像與同行')
    for page in page_list:
        for item in page['items']:
            kind = item['kind']
            if kind == 'rect':
                pdf.setFillColorRGB(*colour(item['fill']))
                pdf.setStrokeColorRGB(*colour(item['stroke'] or item['fill']))
                pdf.roundRect(item['x'], H - item['y'] - item['height'], item['width'], item['height'], item['radius'], stroke=int(bool(item['stroke'])), fill=1)
            elif kind == 'image':
                g = image_geometry(item)
                pdf.saveState()
                path = pdf.beginPath(); path.rect(g['x'], H - g['y'] - g['height'], g['width'], g['height'])
                pdf.clipPath(path, stroke=0)
                pdf.drawImage(pdf_image(item['path']), g['full_x'], H - g['full_y'] - g['full_height'], g['full_width'], g['full_height'], mask='auto')
                pdf.restoreState()
            else:
                pdf.setFillColorRGB(*colour(item['colour']))
                baseline = H - item['y'] - item['ascent']
                obj = pdf.beginText(item['x'], baseline)
                for ch, family, fs in zip(item['text'], item['fonts'], item['sizes']):
                    obj.setFont(family, fs); obj.textOut(ch)
                pdf.drawText(obj)
        pdf.showPage()
    pdf.save()


def render_pptx(page_list, target):
    deck = Presentation(); deck.slide_width = Pt(W); deck.slide_height = Pt(H)
    deck.core_properties.title = BOOK['title']; deck.core_properties.subject = '走部原創繪本・可編輯排版'
    for page in page_list:
        slide = deck.slides.add_slide(deck.slide_layouts[6])
        slide.notes_slide.notes_text_frame.text = page['note']
        for item in page['items']:
            if item['kind'] == 'rect':
                shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if item['radius'] else MSO_SHAPE.RECTANGLE,
                                               Pt(item['x']), Pt(item['y']), Pt(item['width']), Pt(item['height']))
                shape.fill.solid(); shape.fill.fore_color.rgb = RGBColor.from_string(item['fill'][1:])
                if item['stroke']:
                    shape.line.color.rgb = RGBColor.from_string(item['stroke'][1:]); shape.line.width = Pt(.6)
                else:
                    shape.line.fill.background()
            elif item['kind'] == 'image':
                g = image_geometry(item)
                shape = slide.shapes.add_picture(item['path'], Pt(g['x']), Pt(g['y']), Pt(g['width']), Pt(g['height']))
                l, t, r, b = g['crop']
                shape.crop_left = l; shape.crop_top = t; shape.crop_right = 1-r; shape.crop_bottom = 1-b
                shape._element.xpath('.//p:cNvPr')[0].set('descr', item['alt'])
            else:
                shape = slide.shapes.add_textbox(Pt(item['x']), Pt(item['y']), Pt(item['width'] + 5), Pt(item['height'] + 4))
                tf = shape.text_frame; tf.word_wrap = False
                tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
                p = tf.paragraphs[0]; p.space_before = p.space_after = Pt(0); p.line_spacing = 1.0
                for ch, family, fs in zip(item['text'], item['fonts'], item['sizes']):
                    run = p.add_run(); run.text = ch; run.font.name = family; run.font.size = Pt(fs)
                    run.font.bold = family == HAND_FAMILY
                    run.font.color.rgb = RGBColor.from_string(item['colour'][1:])
                    ea = OxmlElement('a:ea'); ea.set('typeface', family); run._r.get_or_add_rPr().append(ea)
    deck.save(str(target))


def render_svgs(page_list, directory):
    for page in page_list:
        elements = []
        for item in page['items']:
            if item['kind'] == 'rect':
                elements.append(f'<rect x="{item["x"]}" y="{item["y"]}" width="{item["width"]}" height="{item["height"]}" rx="{item["radius"]}" fill="{item["fill"]}" stroke="{item["stroke"] or "none"}"/>')
            elif item['kind'] == 'image':
                g = image_geometry(item); iw, ih = g['source']; l, t, r, b = g['crop']
                href = os.path.relpath(item['path'], directory)
                elements.append(f'<svg x="{g["x"]}" y="{g["y"]}" width="{g["width"]}" height="{g["height"]}" viewBox="{l*iw} {t*ih} {(r-l)*iw} {(b-t)*ih}"><title>{html.escape(item["alt"])}</title><image href="{html.escape(href)}" width="{iw}" height="{ih}"/></svg>')
            else:
                baseline = item['y'] + item['ascent']
                spans = ''.join(f'<tspan font-family="{family}" font-weight="{700 if family == HAND_FAMILY else 400}" font-size="{fs}">{html.escape(ch)}</tspan>' for ch, family, fs in zip(item['text'], item['fonts'], item['sizes']))
                elements.append(f'<text x="{item["x"]}" y="{baseline}" fill="{item["colour"]}" font-family="Iansui,芫荽">{spans}</text>')
        content = '<svg xmlns="http://www.w3.org/2000/svg" width="960" height="540" viewBox="0 0 960 540"><rect width="960" height="540" fill="white"/>' + '\n'.join(elements) + '</svg>'
        (directory / (page['id'] + '.svg')).write_text(content)
        ET.fromstring(content)


def single_pages(source, target):
    spread = fitz.open(source); book = fitz.open()
    for index in range(len(spread)):
        if index == 0:
            clips = [fitz.Rect(240, 0, 720, H)]
        elif index == 1:
            clips = [fitz.Rect(W/2, 0, W, H)]
        else:
            clips = [fitz.Rect(0, 0, W/2, H), fitz.Rect(W/2, 0, W, H)]
        for clip in clips:
            page = book.new_page(width=W/2, height=H)
            page.show_pdf_page(page.rect, spread, index, clip=clip)
    book.set_metadata({'title': BOOK['title'] + '｜走部故事', 'author': '點點書房・原創繪本製作', 'subject': '7–9 歲；封面、32內頁、封底'})
    book.set_toc([[1,'封面',1],[1,'故事',3],[1,'走部遊戲',31],[1,'製作頁',33],[1,'封底',34]])
    book.save(target, garbage=4, deflate=True)


def validate(page_list, output, proof):
    cmaps = {family: FontToolsFont(path).getBestCmap() for family, path in FONTS.items()}
    missing = sorted({ch for p in page_list for i in p['items'] if i['kind']=='text' for ch, family in zip(i['text'], i['fonts']) if ord(ch) not in cmaps[family]})
    assert not missing, 'Missing glyphs: ' + ''.join(missing)
    for p in page_list:
        for i in p['items']:
            if i['kind'] == 'text':
                assert i['x'] >= 0 and i['x'] + i['width'] <= W + .1 and i['y'] + i['height'] <= H, (p['id'], i['text'], 'outside page')
                if p['id'] != 'cover':
                    assert not (i['x'] < W/2 and i['x'] + i['width'] > W/2), (p['id'], i['text'], 'text crosses binding')
                else:
                    assert 240 <= i['x'] and i['x'] + i['width'] <= 720, (p['id'], 'text outside cover trim')
    spreads = fitz.open(output / '走！去找那朵雲_跨頁閱讀版.pdf')
    leaves = fitz.open(output / '走！去找那朵雲_繪本.pdf')
    assert len(spreads) == 18 and len(leaves) == 34
    leaf_index = 0
    for index, page in enumerate(page_list):
        if index == 0:
            clips = [(240,720)]
        elif index == 1:
            clips = [(480,960)]
        else:
            clips = [(0,480),(480,960)]
        for left,right in clips:
            expected = ''.join(i['text'] for i in page['items'] if i['kind']=='text' and i['x'] >= left and i['x'] + i['width'] <= right)
            actual = leaves[leaf_index].get_text(clip=leaves[leaf_index].rect)
            assert Counter(re.sub(r'\s','',expected)) == Counter(re.sub(r'\s','',actual)), ('Clipped or missing text',leaf_index+1)
            leaf_index += 1
    full_text = ''.join(p.get_text() for p in spreads)
    for ch in BOOK['coreCharacters'] + BOOK['extensionCharacters'] + [BOOK['radical']['base']]:
        assert ch in full_text, 'PDF text missing: ' + ch
    deck = Presentation(output / '走！去找那朵雲_可編輯排版.pptx')
    assert len(deck.slides) == 18
    assert sum(len(s.shapes) for s in deck.slides) > 100
    expected = ''.join(i['text'] for p in page_list for i in p['items'] if i['kind']=='text')
    actual = ''.join(shape.text for slide in deck.slides for shape in slide.shapes if shape.has_text_frame)
    assert expected == actual, 'Editable slide text differs from book source'
    # Verify the exported runs: fallback or flattened text would defeat the learning cue.
    handwritten = Counter()
    for layout, pdf_page, slide in zip(page_list, spreads, deck.slides):
        text_items = [item for item in layout['items'] if item['kind'] == 'text']
        expected_runs = [(ch, family, fs) for item in text_items
                         for ch, family, fs in zip(item['text'], item['fonts'], item['sizes'])]
        actual_runs = [(run.text, run.font.name, run.font.size.pt)
                       for shape in slide.shapes if shape.has_text_frame
                       for para in shape.text_frame.paragraphs for run in para.runs]
        assert len(actual_runs) == len(expected_runs), (layout['id'], 'Editable font runs missing')
        for (ch, family, fs), (text, exported_font, exported_size) in zip(expected_runs, actual_runs):
            assert ch == text and family == exported_font and abs(fs - exported_size) < .011, (layout['id'], ch, 'Editable font or size changed')
        expected_hand = Counter((ch, round(fs, 2)) for ch, family, fs in expected_runs if family == HAND_FAMILY)
        pdf_hand = Counter((char['c'], round(span['size'], 2))
                           for block in pdf_page.get_text('rawdict')['blocks'] if 'lines' in block
                           for line in block['lines'] for span in line['spans'] if HAND_FAMILY in span['font']
                           for char in span['chars'])
        assert pdf_hand == expected_hand, (layout['id'], 'PDF handwriting font or size changed')
        handwritten.update(ch for ch, family, fs in expected_runs if family == HAND_FAMILY)
        if layout['id'].startswith('s'):
            for item in text_items:
                for ch, family, fs in zip(item['text'], item['fonts'], item['sizes']):
                    if ch in TARGETS:
                        assert family == HAND_FAMILY and abs(fs / item['size'] - EMPHASIS_SCALE) < .001
    report = {'status': 'proof' if proof else 'complete', 'spreadPages':len(spreads), 'singlePages':len(leaves),
              'editableSlides':len(deck.slides), 'missingGlyphs':missing, 'allTextEditable':True,
              'targetTypography': {'font': HAND_FAMILY, 'bodyScale': EMPHASIS_SCALE,
                                   'pdfAndPptxRunsVerified': True, 'handwrittenGlyphs': dict(handwritten)},
              'storyHanCharacters':len(re.findall(r'[\u3400-\u9fff]', ''.join(''.join(s['lines']) for s in BOOK['spreads']))),
              'targets':{ch:full_text.count(ch) for ch in BOOK['coreCharacters']+BOOK['extensionCharacters']}}
    (output / 'validation.json').write_text(json.dumps(report,ensure_ascii=False,indent=2)+'\n')
    return report


def main():
    parser = argparse.ArgumentParser(description=__doc__); parser.add_argument('--proof', action='store_true')
    args = parser.parse_args()
    output = ROOT / ('.cache/proof' if args.proof else 'exports')
    editable = output / 'editable' if args.proof else ROOT / 'editable'
    output.mkdir(parents=True, exist_ok=True); editable.mkdir(parents=True, exist_ok=True)
    page_list = pages(args.proof)
    render_pdf(page_list, output / '走！去找那朵雲_跨頁閱讀版.pdf')
    single_pages(output / '走！去找那朵雲_跨頁閱讀版.pdf', output / '走！去找那朵雲_繪本.pdf')
    render_pptx(page_list, output / '走！去找那朵雲_可編輯排版.pptx')
    render_svgs(page_list, editable)
    report = validate(page_list, output, args.proof)
    doc = fitz.open(output / '走！去找那朵雲_跨頁閱讀版.pdf')
    preview = ROOT / '.cache' / ('proof/previews' if args.proof else 'final/previews')
    preview.mkdir(parents=True, exist_ok=True)
    thumbs = []
    for index, page in enumerate(doc):
        pix = page.get_pixmap(matrix=fitz.Matrix(1,1), alpha=False)
        pix.save(preview / f'{index+1:02}.png')
        thumb = Image.frombytes('RGB',[pix.width,pix.height],pix.samples); thumb.thumbnail((320,180))
        tile=Image.new('RGB',(336,208),'#e9ece9');tile.paste(thumb,(8,18));ImageDraw.Draw(tile).text((8,3),str(index+1),fill='#303536');thumbs.append(tile)
    sheet=Image.new('RGB',(336*3,208*6),'white')
    for index,thumb in enumerate(thumbs):sheet.paste(thumb,((index%3)*336,(index//3)*208))
    sheet.save(output / '全書縮圖.png')
    print(json.dumps(report,ensure_ascii=False,indent=2))


if __name__ == '__main__':
    main()
